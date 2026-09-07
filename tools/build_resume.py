from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(__file__).resolve().parents[1] / "resume" / "William_Xie_Manager_Data_Analytics_AI_Resume.docx"

INK = RGBColor(18, 35, 45)
ACCENT = RGBColor(196, 145, 2)
MUTED = RGBColor(88, 101, 109)
LIGHT = "F3F5F6"
GOLD_LIGHT = "FBF5DF"
WHITE = RGBColor(255, 255, 255)


def set_run(run, size=None, color=INK, bold=None, italic=None, font="Aptos"):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    if size is not None:
        run.font.size = Pt(size)
    run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent=0):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[i]))
            tc_w.set(qn("w:type"), "dxa")


def add_bottom_border(paragraph, color="C49102", size="12", space="4"):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    borders.append(bottom)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = paragraph.add_run("Page ")
    set_run(r, size=8.5, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    r2 = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "58656D")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "17")
    rpr.extend([color, size])
    r2.append(rpr)
    t = OxmlElement("w:t")
    t.text = "1"
    r2.append(t)
    fld.append(r2)
    paragraph._p.append(fld)


def make_bullet_numbering(doc):
    numbering = doc.part.numbering_part.element
    abs_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(abs_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet")
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•")
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    p_pr.extend([tabs, ind])
    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Arial")
    fonts.set(qn("w:hAnsi"), "Arial")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "C49102")
    r_pr.extend([fonts, color])
    lvl.extend([start, num_fmt, lvl_text, lvl_jc, p_pr, r_pr])
    abstract.append(lvl)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abs_ref = OxmlElement("w:abstractNumId")
    abs_ref.set(qn("w:val"), str(abstract_id))
    num.append(abs_ref)
    numbering.append(num)
    return num_id


def add_bullet(doc, num_id, text):
    p = doc.add_paragraph(style="Resume Bullet")
    p_pr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    numid = OxmlElement("w:numId")
    numid.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, numid])
    p_pr.append(num_pr)
    r = p.add_run(text)
    set_run(r, size=9.6)
    return p


def add_section_heading(doc, text):
    p = doc.add_paragraph(style="Heading 1")
    r = p.add_run(text.upper())
    set_run(r, size=11.5, color=INK, bold=True)
    add_bottom_border(p, color="C49102", size="8", space="3")
    return p


def add_role(doc, company, role, dates):
    p = doc.add_paragraph(style="Role Header")
    r = p.add_run(company)
    set_run(r, size=10.6, color=INK, bold=True)
    r = p.add_run("  |  ")
    set_run(r, size=10.2, color=ACCENT, bold=True)
    r = p.add_run(role)
    set_run(r, size=10.2, color=INK, bold=True)
    r = p.add_run(f"\n{dates}")
    set_run(r, size=8.8, color=MUTED, italic=True)


def add_project(doc, title, text):
    p = doc.add_paragraph(style="Project")
    r = p.add_run(title + ": ")
    set_run(r, size=9.6, bold=True, color=INK)
    r = p.add_run(text)
    set_run(r, size=9.6, color=INK)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.62)
section.bottom_margin = Inches(0.58)
section.left_margin = Inches(0.72)
section.right_margin = Inches(0.72)
section.header_distance = Inches(0.28)
section.footer_distance = Inches(0.28)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
normal.font.size = Pt(9.8)
normal.font.color.rgb = INK
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(4)
normal.paragraph_format.line_spacing = 1.12

h1 = styles["Heading 1"]
h1.font.name = "Aptos"
h1.font.size = Pt(11.5)
h1.font.bold = True
h1.font.color.rgb = INK
h1.paragraph_format.space_before = Pt(10)
h1.paragraph_format.space_after = Pt(6)
h1.paragraph_format.keep_with_next = True

for name in ("Resume Bullet", "Role Header", "Project"):
    if name not in styles:
        styles.add_style(name, 1)

bullet_style = styles["Resume Bullet"]
bullet_style.base_style = normal
bullet_style.font.name = "Aptos"
bullet_style.font.size = Pt(9.6)
bullet_style.paragraph_format.space_after = Pt(2.8)
bullet_style.paragraph_format.line_spacing = 1.09

role_style = styles["Role Header"]
role_style.base_style = normal
role_style.paragraph_format.space_before = Pt(6)
role_style.paragraph_format.space_after = Pt(3)
role_style.paragraph_format.keep_with_next = True

project_style = styles["Project"]
project_style.base_style = normal
project_style.paragraph_format.space_after = Pt(5)
project_style.paragraph_format.line_spacing = 1.1

header = section.header
hp = header.paragraphs[0]
hp.text = "WILLIAM XIE  |  DATA ANALYTICS & AI"
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
set_run(hp.runs[0], size=8, color=MUTED, bold=True)
footer = section.footer
add_page_number(footer.paragraphs[0])

# Memo-masthead-inspired identity block.
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(0)
r = p.add_run("WILLIAM XIE")
set_run(r, size=26, color=INK, bold=True)
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(0)
p.paragraph_format.space_after = Pt(4)
r = p.add_run("DATA ANALYTICS & AI LEADER")
set_run(r, size=12.5, color=ACCENT, bold=True)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(10)
r = p.add_run("Sydney, Australia  |  [PHONE]  |  [EMAIL]  |  linkedin.com/in/williamzmx")
set_run(r, size=9, color=MUTED)
add_bottom_border(p, color="C49102", size="12", space="4")

add_section_heading(doc, "Executive Profile")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(7)
r = p.add_run(
    "Commercially focused analytics and AI leader with 6+ years of consistently achieving performance objectives and a record of translating ambiguous business questions into decision-ready insights, scalable reporting, automation and machine-learning solutions. Combines hands-on SQL, Python, Power BI/Tableau and Azure delivery with stakeholder confidence across Finance, Category, Marketing, Sales, IT and Operations. Led pricing and competitive-intelligence initiatives that increased profit by 3%+ across a $500M+ portfolio, reduced workload by more than 50% through automation, and helped deliver 20+ cross-functional projects."
)
set_run(r, size=9.8)

metrics = doc.add_table(rows=1, cols=4)
set_table_geometry(metrics, [2556, 2556, 2556, 2556], indent=0)
for cell in metrics.rows[0].cells:
    shade_cell(cell, GOLD_LIGHT)
    set_cell_margins(cell, top=110, bottom=110, start=110, end=110)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
metric_values = [("3%+", "profit uplift"), ("$500M+", "portfolio influenced"), ("50%+", "effort saved"), ("20+", "projects delivered")]
for cell, (value, label) in zip(metrics.rows[0].cells, metric_values):
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(value + "\n")
    set_run(r, size=14, color=INK, bold=True)
    r = p.add_run(label.upper())
    set_run(r, size=7.2, color=MUTED, bold=True)

add_section_heading(doc, "Core Expertise")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(2)
r = p.add_run(
    "Finance & commercial analytics  •  Executive reporting & data storytelling  •  AI use-case discovery  •  Data quality & governance  •  Pricing & profitability  •  Dashboard product design  •  Stakeholder consulting  •  Process automation  •  Forecasting & machine learning  •  Team leadership & mentoring"
)
set_run(r, size=9.2, color=INK, bold=True)

bullet_num_id = make_bullet_numbering(doc)

add_section_heading(doc, "Professional Experience")
add_role(doc, "[CURRENT EMPLOYER]", "[CURRENT TITLE - Analytics / Pricing / AI]", "Sydney, Australia  |  [MM/YYYY - Present]")
add_bullet(doc, bullet_num_id, "Led a cross-functional price-change program using elasticity modelling, competitive-pricing evidence, customer feedback and structured testing; increased profit by 3%+ across a portfolio exceeding AUD $500M annually.")
add_bullet(doc, bullet_num_id, "Translated commercial and finance questions into reusable analytics assets, including product clustering, a centralised pricing reference, price-evaluation models and web-scraped market intelligence.")
add_bullet(doc, bullet_num_id, "Designed cloud-based automation and data workflows using Azure Databricks, Data Factory, DevOps, Blob Storage, Python, PySpark, SQL and REST/SOAP APIs, reducing manual effort and improving process reliability.")
add_bullet(doc, bullet_num_id, "Partnered with Finance, Category, Marketing, Sales, IT and Operations to define problems, test recommendations, improve pricing logic and customer experience, and move strategic initiatives into delivery.")
add_bullet(doc, bullet_num_id, "Shaped applied AI opportunities including Azure ML-based dynamic pricing and a real-time recommender system for cross-sell and up-sell use cases; managed SQL Server Agent jobs and BAU data controls.")
add_bullet(doc, bullet_num_id, "Supported ERP adoption, app logic and UX improvements through testing, documentation, knowledge sharing and structured stakeholder feedback.")

add_role(doc, "[PREVIOUS EMPLOYER]", "[PREVIOUS TITLE - Data Analyst / Data Scientist]", "[LOCATION]  |  [MM/YYYY - MM/YYYY]")
add_bullet(doc, bullet_num_id, "Built and maintained 100+ external and internal reports and decision products using Power BI, Tableau and statistical analysis, earning positive stakeholder feedback.")
add_bullet(doc, bullet_num_id, "Created and managed SQL Server and PostgreSQL databases, data pipelines and migrations; used SSIS, Python and web-crawling methods to deliver reliable reporting data.")
add_bullet(doc, bullet_num_id, "Developed CNN, NLP/sentiment, forecasting and LSTM models to improve design relevance, marketing content, customer understanding, market-inventory estimates and commodity-price forecasts.")
add_bullet(doc, bullet_num_id, "Automated recurring ETL, reporting, website, invoicing, accounts-receivable and customer-engagement processes using Python, Google Apps Script, APIs and virtual machines, saving more than 50% of required workload in selected processes.")
continuation = doc.add_paragraph(style="Role Header")
continuation.paragraph_format.page_break_before = True
continuation.paragraph_format.space_after = Pt(3)
r = continuation.add_run("[PREVIOUS EMPLOYER]  |  CONTINUED")
set_run(r, size=9.2, color=MUTED, bold=True)
add_bullet(doc, bullet_num_id, "Managed 550+ email and digital campaigns, using segmentation, A/B testing, send-time optimisation and content analysis to achieve open and click rates more than 100% above industry benchmarks.")
add_bullet(doc, bullet_num_id, "Led and supported cross-functional delivery, trained junior analysts, presented evidence to executives and customers, and contributed to winning a major European buyer opportunity through forecasting and interactive Power BI storytelling.")

add_section_heading(doc, "Selected Analytics & AI Portfolio")
add_project(doc, "Pricing, profitability and competitive intelligence", "Designed elasticity and price-evaluation models, a product-clustering capability and large-scale web scraping to identify under- and over-priced products, support tender selection and improve portfolio profitability.")
add_project(doc, "AI-enabled customer and product decisions", "Built computer-vision models for design preference, NLP/sentiment models for customer interest, forecasting models for cashflow and commodities, and recommender concepts for cross-sell and up-sell opportunities.")
add_project(doc, "Scalable reporting and decision products", "Developed and maintained Power BI/Tableau reporting products, geospatial views, market-inventory forecasts and executive narratives that translated complex evidence into clear actions.")
add_project(doc, "Automation and data assets", "Created repeatable ETL and API workflows across Azure and SQL platforms, improved data synchronisation and BAU processing, and reduced manual effort, error risk and delivery time.")

add_section_heading(doc, "Leadership & Stakeholder Value")
add_bullet(doc, bullet_num_id, "Consultative problem solving: converts broad stakeholder questions into clear problem statements, practical analysis and evidence-linked recommendations.")
add_bullet(doc, bullet_num_id, "Executive communication: communicates technical findings to Finance, company boards, operational leaders, customers and non-technical stakeholders through concise reports, dashboards and presentations.")
add_bullet(doc, bullet_num_id, "Delivery leadership: contributed as a lead or principal contributor across 20+ projects, supported team development and maintained a consistent record of achieving KPIs and agreed deliverables.")
add_bullet(doc, bullet_num_id, "Commercial curiosity: connects customer, product, channel, pricing and operational data to revenue, profit, engagement and efficiency outcomes.")

add_section_heading(doc, "Technology")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(3)
r = p.add_run("Analytics & AI: ")
set_run(r, size=9.4, bold=True)
r = p.add_run("Python, R, PySpark, scikit-learn, TensorFlow, Azure ML Studio, forecasting, NLP, computer vision, recommender systems, A/B testing")
set_run(r, size=9.4)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(3)
r = p.add_run("Data & cloud: ")
set_run(r, size=9.4, bold=True)
r = p.add_run("T-SQL, PostgreSQL, P/SQL, SQL Server Agent, SSIS, Azure Databricks, Data Factory, Blob Storage, Azure DevOps, Apache Spark, REST/SOAP APIs, Amazon EC2/S3")
set_run(r, size=9.4)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(3)
r = p.add_run("Reporting & visualisation: ")
set_run(r, size=9.4, bold=True)
r = p.add_run("Power BI, Tableau, matplotlib, folium, geospatial analysis, executive insight packs, dashboard design and data storytelling")
set_run(r, size=9.4)

add_section_heading(doc, "Education & Credentials")
p = doc.add_paragraph()
r = p.add_run("[DEGREE / QUALIFICATION]  |  [INSTITUTION]  |  [YEAR]")
set_run(r, size=9.6, bold=True)
p = doc.add_paragraph()
r = p.add_run("[RELEVANT CERTIFICATIONS, PROFESSIONAL DEVELOPMENT OR BANKING / AI GOVERNANCE TRAINING]")
set_run(r, size=9.2, color=MUTED)

add_section_heading(doc, "Additional Achievement")
p = doc.add_paragraph()
r = p.add_run("Earlier sales experience included breaking a company sales record and leading a high-performing sales team, strengthening William's customer focus, commercial judgement and confidence in stakeholder-facing environments.")
set_run(r, size=9.5)

# Keep section headings with their first content paragraph and avoid split role headers.
for p in doc.paragraphs:
    if p.style.name in ("Heading 1", "Role Header"):
        p.paragraph_format.keep_with_next = True

doc.core_properties.title = "William Xie - Manager Data Analytics and AI Resume"
doc.core_properties.subject = "Targeted resume for Manager Data Analytics and AI"
doc.core_properties.author = "William Xie"
doc.core_properties.keywords = "data analytics, AI, SQL, Python, Power BI, finance, governance, reporting"

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
