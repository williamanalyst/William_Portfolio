from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(__file__).resolve().parents[1] / "resume" / "William_Xie_Lead_Data_Analyst_Resume.docx"
INK = RGBColor(0, 0, 0)
GRAY = RGBColor(89, 89, 89)


def set_font(run, size, bold=False, color=INK, italic=False):
    run.font.name = "Aptos"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Aptos")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold = bold
    run.italic = italic


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run("William Xie | Lead Data Analyst | ")
    set_font(run, 8, color=GRAY)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run_element = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "16")
    props.append(size)
    run_element.append(props)
    text = OxmlElement("w:t")
    text.text = "1"
    run_element.append(text)
    field.append(run_element)
    paragraph._p.append(field)


def add_heading(doc, text):
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text.upper())
    set_font(r, 10.5, bold=True)
    return p


def add_role(doc, employer, title, location_dates):
    p = doc.add_paragraph(style="Role")
    p.paragraph_format.keep_with_next = True
    r = p.add_run(employer)
    set_font(r, 10.3, bold=True)
    r = p.add_run(" | " + title)
    set_font(r, 10.3, bold=True)
    p = doc.add_paragraph(style="Role Detail")
    p.paragraph_format.keep_with_next = True
    r = p.add_run(location_dates)
    set_font(r, 9, color=GRAY, italic=True)


def add_bullet(doc, text):
    p = doc.add_paragraph(style="Resume Bullet")
    p.paragraph_format.keep_together = True
    r = p.add_run(text)
    set_font(r, 9.35)


def add_labeled_paragraph(doc, label, text):
    p = doc.add_paragraph(style="Compact")
    r = p.add_run(label + ": ")
    set_font(r, 9.25, bold=True)
    r = p.add_run(text)
    set_font(r, 9.25)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.54)
section.bottom_margin = Inches(0.5)
section.left_margin = Inches(0.7)
section.right_margin = Inches(0.7)
section.header_distance = Inches(0.25)
section.footer_distance = Inches(0.25)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
normal.font.size = Pt(9.35)
normal.font.color.rgb = INK
normal.paragraph_format.space_after = Pt(3)
normal.paragraph_format.line_spacing = 1.07

title = styles["Title"]
title.font.name = "Aptos"
title.font.size = Pt(22)
title.font.bold = True
title.font.color.rgb = INK
title.paragraph_format.space_after = Pt(2)

h1 = styles["Heading 1"]
h1.font.name = "Aptos"
h1.font.size = Pt(10.5)
h1.font.bold = True
h1.font.color.rgb = INK
h1.paragraph_format.space_before = Pt(8)
h1.paragraph_format.space_after = Pt(3)

for name in ("Role", "Role Detail", "Compact", "Resume Bullet"):
    if name not in styles:
        styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)

role = styles["Role"]
role.base_style = normal
role.paragraph_format.space_before = Pt(5)
role.paragraph_format.space_after = Pt(0)

role_detail = styles["Role Detail"]
role_detail.base_style = normal
role_detail.paragraph_format.space_after = Pt(2)

compact = styles["Compact"]
compact.base_style = normal
compact.paragraph_format.space_after = Pt(2)
compact.paragraph_format.line_spacing = 1.04

bullet = styles["Resume Bullet"]
bullet.base_style = normal
bullet.paragraph_format.left_indent = Inches(0.17)
bullet.paragraph_format.first_line_indent = Inches(-0.12)
bullet.paragraph_format.space_after = Pt(1.7)
bullet.paragraph_format.line_spacing = 1.04

footer = section.footer.paragraphs[0]
add_page_number(footer)

p = doc.add_paragraph(style="Title")
r = p.add_run("William Xie Lead Data Analyst Resume")
set_font(r, 22, bold=True)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(5)
r = p.add_run("Sydney, Australia | [Add phone] | [Add email] | linkedin.com/in/williamzmx")
set_font(r, 9.3, color=GRAY)

add_heading(doc, "Lead Data Analyst Profile")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(4)
r = p.add_run(
    "Commercially focused analytics leader with 6+ years of experience translating ambiguous business questions into decision-ready insights, scalable reporting, automation and applied machine-learning solutions. Combines hands-on SQL, Python, Power BI, Tableau and Azure delivery with the stakeholder leadership to turn analysis into action. Delivered 3%+ profit uplift across a $500M+ annual portfolio, reduced effort by more than 50% in selected workflows, and contributed to 20+ cross-functional initiatives."
)
set_font(r, 9.45)

add_heading(doc, "Leadership and Technical Strengths")
add_labeled_paragraph(doc, "Lead analytics", "Commercial and financial analysis, pricing and profitability, KPI design, executive reporting, dashboard product ownership, data quality and governance")
add_labeled_paragraph(doc, "Influence and delivery", "Stakeholder discovery, requirements translation, cross-functional delivery, experimentation, insight storytelling, mentoring and knowledge sharing")
add_labeled_paragraph(doc, "Data and BI", "SQL, Python, Power BI, Tableau, T-SQL, PostgreSQL, SSIS, Azure Databricks, Data Factory, Azure DevOps, PySpark, REST and SOAP APIs")
add_labeled_paragraph(doc, "Advanced analytics", "Price elasticity, product clustering, forecasting, A/B testing, NLP and sentiment analysis, computer vision, recommender systems, scikit-learn, TensorFlow and Azure ML")

add_heading(doc, "Professional Experience")
add_role(doc, "[Current employer]", "[Current title - Analytics, Pricing and AI]", "Sydney, Australia | [MM/YYYY] - Present")
add_bullet(doc, "Led a cross-functional price-change program using elasticity modelling, competitor evidence, customer feedback and structured testing, increasing profit by 3%+ across a portfolio exceeding AUD $500M annually.")
add_bullet(doc, "Turned commercial and finance questions into reusable decision assets, including a central pricing reference, product clustering, price-evaluation models and web-scraped market intelligence.")
add_bullet(doc, "Designed and delivered automated data workflows using Azure Databricks, Data Factory, Blob Storage, Python, PySpark, SQL and APIs, improving reliability and reducing manual processing.")
add_bullet(doc, "Partnered with Finance, Category, Marketing, Sales, IT and Operations to frame problems, validate recommendations and embed new pricing logic, reporting and customer-experience improvements.")
add_bullet(doc, "Led applied AI discovery for dynamic pricing and real-time cross-sell and up-sell recommendations; maintained SQL Server Agent jobs and BAU data controls to support trusted operations.")
add_bullet(doc, "Supported ERP adoption and product improvements through testing, documentation, stakeholder feedback and practical knowledge sharing.")

add_role(doc, "[Previous employer]", "[Data Analyst / Data Scientist]", "[Location] | [MM/YYYY] - [MM/YYYY]")
add_bullet(doc, "Built and maintained 100+ internal and external reports and decision products in Power BI and Tableau, converting complex data into accessible recommendations for stakeholders.")
add_bullet(doc, "Created and managed SQL Server and PostgreSQL databases, data pipelines and migrations, using SSIS, Python and web-crawling approaches to supply reliable reporting data.")
add_bullet(doc, "Automated recurring ETL, reporting, invoicing, accounts-receivable and customer-engagement processes using Python, Google Apps Script, APIs and virtual machines, saving more than 50% of effort in selected workflows.")

p = doc.add_paragraph()
p.paragraph_format.page_break_before = True
p.paragraph_format.space_after = Pt(0)

add_heading(doc, "Professional Experience Continued")
add_role(doc, "[Previous employer]", "[Data Analyst / Data Scientist]", "[Location] | [MM/YYYY] - [MM/YYYY]")
add_bullet(doc, "Developed forecasting, CNN and NLP/sentiment models to improve design relevance, marketing content, customer understanding, market-inventory estimates and commodity-price forecasts.")
add_bullet(doc, "Managed 550+ email and digital campaigns using segmentation, A/B testing, send-time optimisation and content analysis, achieving open and click rates more than 100% above industry benchmarks.")
add_bullet(doc, "Presented evidence to executives and customers, trained junior analysts and supported cross-functional delivery; contributed to winning a major European buyer opportunity through forecasting and interactive Power BI storytelling.")

add_heading(doc, "Selected Leadership Impact")
add_labeled_paragraph(doc, "Commercial value", "Connected pricing, customer, product and market signals to profitability decisions; identified opportunities at portfolio scale and created evidence stakeholders could test and act on.")
add_labeled_paragraph(doc, "Decision products", "Built Power BI, Tableau, geospatial and executive reporting products that made complex analysis useful to senior, operational and non-technical audiences.")
add_labeled_paragraph(doc, "Reliable delivery", "Combined model development with data pipelines, BAU controls, testing, documentation and adoption support so analytics could operate beyond a proof of concept.")
add_labeled_paragraph(doc, "People leadership", "Worked as a lead contributor across 20+ initiatives, coached junior analysts and created shared ways of working across commercial and technical teams.")

add_heading(doc, "Education and Credentials")
p = doc.add_paragraph(style="Compact")
r = p.add_run("[Add degree or qualification] | [Institution] | [Year]")
set_font(r, 9.25, bold=True)
p = doc.add_paragraph(style="Compact")
r = p.add_run("[Add relevant certifications, professional development or AI and data governance training]")
set_font(r, 9.25, color=GRAY)

add_heading(doc, "Additional Commercial Experience")
p = doc.add_paragraph()
r = p.add_run("Earlier sales leadership included breaking a company sales record and leading a high-performing team, strengthening commercial judgement, customer focus and stakeholder confidence.")
set_font(r, 9.35)

doc.core_properties.title = "William Xie Lead Data Analyst Resume"
doc.core_properties.subject = "Lead Data Analyst resume tailored for Sydney opportunities"
doc.core_properties.author = "William Xie"
doc.core_properties.keywords = "lead data analyst, Sydney, SQL, Python, Power BI, Azure, commercial analytics"
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
